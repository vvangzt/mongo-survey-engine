"""Build the polished second-stage report DOCX from docs/project-report.md.

The Markdown file is the content source of truth. This renderer owns the
academic layout: R5 cover, front-matter TOC, Arabic body page numbers,
three-line tables, code blocks, and embedded screenshots.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "project-report.md"
TARGET = ROOT / "docs" / "项目完成报告.docx"
SCREENSHOT_DIR = ROOT / "docs"

BODY_CN = "SimSun"
BODY_EN = "Times New Roman"
HEAD_CN = "SimHei"
MONO = "Courier New"

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x60, 0x60, 0x60)
LIGHT_GRAY = RGBColor(0x80, 0x80, 0x80)
CODE_COLOR = RGBColor(0x1F, 0x1F, 0x1F)

CONTENT_WIDTH_IN = 6.0
BODY_SIZE = 12
CAPTION_SIZE = 10.5
TABLE_SIZE = 9
CODE_SIZE = 8.5
LINE = 1.3


# ---------------------------------------------------------------------------
# Low-level OOXML helpers
# ---------------------------------------------------------------------------


def set_run_font(
    run,
    *,
    cn: str = BODY_CN,
    en: str = BODY_EN,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:cs"), en)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, *, before: float = 0, after: float = 3, line: float = LINE):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_bottom_border(paragraph, size: int = 6, color: str = "BFBFBF") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def shade_paragraph(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_paragraph_border(paragraph, size: int = 4, color: str = "D9D9D9") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), color)
        pbdr.append(el)


def set_cell_border(cell, **kwargs) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        data = kwargs[edge]
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        if data is None:
            element.set(qn("w:val"), "nil")
            continue
        element.set(qn("w:val"), data.get("val", "single"))
        element.set(qn("w:sz"), str(data.get("sz", 4)))
        element.set(qn("w:space"), str(data.get("space", 0)))
        element.set(qn("w:color"), data.get("color", "000000"))


def set_cell_margins(cell, *, top=0, bottom=0, left=0, right=0) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(int(value.twips if hasattr(value, "twips") else value)))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def clear_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def set_table_width(table, width) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(width.twips)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)


def set_table_layout_fixed(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def set_page_numbering(section, fmt: str) -> None:
    sect_pr = section._sectPr
    for existing in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(existing)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), "1")
    sect_pr.append(pg)


def add_update_fields_setting(document) -> None:
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_EN
    normal.font.size = Pt(BODY_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    normal.paragraph_format.line_spacing = LINE
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.first_line_indent = Pt(BODY_SIZE * 2)

    heading_specs = {
        "Heading 1": (16, True, WD_ALIGN_PARAGRAPH.CENTER, 18, 10),
        "Heading 2": (15, True, WD_ALIGN_PARAGRAPH.LEFT, 14, 6),
        "Heading 3": (14, True, WD_ALIGN_PARAGRAPH.LEFT, 10, 4),
    }
    for name, (size, bold, align, before, after) in heading_specs.items():
        style = document.styles[name]
        style.font.name = BODY_EN
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = BLACK
        style.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_CN)
        fmt = style.paragraph_format
        fmt.alignment = align
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = LINE
        fmt.keep_with_next = True
        fmt.first_line_indent = Pt(0)

    caption = document.styles["Caption"]
    caption.font.name = BODY_EN
    caption.font.size = Pt(CAPTION_SIZE)
    caption.font.color.rgb = GRAY
    caption.font.italic = False
    caption.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing = LINE
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.first_line_indent = Pt(0)


# ---------------------------------------------------------------------------
# Inline parsing
# ---------------------------------------------------------------------------


INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def strip_markup(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text


def add_inline_runs(
    paragraph,
    text: str,
    *,
    size: float | None = None,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    for token in INLINE_TOKEN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            add_inline_runs(paragraph, token[2:-2], size=size, bold=True, color=color)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, cn=BODY_CN, en=MONO, size=(size - 1 if size else 10.5), bold=bold, color=color or CODE_COLOR)
        elif token.startswith("[") and "](" in token:
            label = token[1 : token.index("](")]
            run = paragraph.add_run(label)
            set_run_font(run, size=size, bold=bold, color=color)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, bold=bold, color=color)


# ---------------------------------------------------------------------------
# Block renderers
# ---------------------------------------------------------------------------


def add_body_paragraph(document, text: str, *, indent: bool = True, caption: bool = False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(paragraph, after=4)
    if caption:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
        add_inline_runs(paragraph, text, size=CAPTION_SIZE, bold=True)
    else:
        paragraph.paragraph_format.first_line_indent = Pt(BODY_SIZE * 2) if indent else Pt(0)
        add_inline_runs(paragraph, text, size=BODY_SIZE)
    return paragraph


def add_list_item(document, text: str, *, ordered: bool) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, after=2)
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(-12)
    if not ordered:
        text = "• " + text
    add_inline_runs(paragraph, text, size=BODY_SIZE)


def add_heading(document, text: str, level: int):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    if text.startswith("附录") or text == "1. 项目概述":
        paragraph.paragraph_format.page_break_before = True
    run = paragraph.add_run(text)
    sizes = {1: 16, 2: 15, 3: 14}
    set_run_font(run, cn=HEAD_CN, size=sizes[level], bold=True, color=BLACK)
    return paragraph


def add_code_block(document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.45)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_together = True
    shade_paragraph(paragraph, "F5F5F5")
    add_paragraph_border(paragraph)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, cn=BODY_CN, en=MONO, size=CODE_SIZE, color=CODE_COLOR)
        if index != len(lines) - 1:
            run.add_break()


def add_blockquote(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Cm(0.8)
    paragraph.paragraph_format.right_indent = Cm(0.5)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(paragraph, after=4)
    add_bottom_border(paragraph, size=4, color="D0D0D0")
    add_inline_runs(paragraph, text, size=BODY_SIZE, color=GRAY)


def parse_table_row(line: str) -> list[str]:
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [cell.strip() for cell in inner.split("|")]


def is_separator(line: str) -> bool:
    stripped = line.strip()
    return bool(
        stripped.startswith("|")
        and set(stripped.replace("|", "").replace(" ", "")) <= {"-", ":"}
    )


def count_table_rows(lines: list[str], start: int) -> int:
    index = start
    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("|"):
            break
        if stripped.startswith(("```", "#", "![", "- ", "1. ", "2. ", "3. ")):
            return 0
        index += 1
    count = 0
    while index < len(lines) and lines[index].strip().startswith("|"):
        if not is_separator(lines[index]):
            count += 1
        index += 1
    return count


def text_weight(text: str) -> int:
    return sum(2 if ord(char) > 127 else 1 for char in text)


def compute_column_widths(rows: list[list[str]]) -> list[float]:
    cols = max(len(row) for row in rows)
    weights = []
    for index in range(cols):
        longest = max(
            (text_weight(strip_markup(row[index])) for row in rows if index < len(row)),
            default=4,
        )
        weights.append(max(8, min(60, longest)))
    minimum = 0.8
    remaining = CONTENT_WIDTH_IN - minimum * cols
    if remaining <= 0:
        return [CONTENT_WIDTH_IN / cols] * cols
    total = sum(weights)
    widths = [minimum + remaining * weight / total for weight in weights]
    if widths[0] < 1.0:
        delta = 1.0 - widths[0]
        widths[0] = 1.0
        rest_total = sum(widths[1:])
        if rest_total > 0:
            for index in range(1, len(widths)):
                widths[index] -= delta * widths[index] / rest_total
    return widths


def add_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    normalized = [row + [""] * (cols - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table)
    set_table_width(table, Inches(CONTENT_WIDTH_IN))
    clear_table_borders(table)

    widths = compute_column_widths(normalized)
    keep_whole_table = len(normalized) <= 6
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for row_index, row in enumerate(normalized):
        row_obj = table.rows[row_index]
        row_obj.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        tr_pr = row_obj._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            tr_pr.append(header)
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(
                cell,
                top=Cm(0.12),
                bottom=Cm(0.12),
                left=Cm(0.14),
                right=Cm(0.14),
            )
            for edge in ("top", "left", "bottom", "right"):
                set_cell_border(cell, **{edge: None})
            if row_index == 0:
                set_cell_border(cell, bottom={"sz": 8, "color": "000000"})
                set_cell_shading(cell, "F2F2F2")
            if row_index == len(normalized) - 1:
                set_cell_border(cell, bottom={"sz": 12, "color": "000000"})
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = LINE
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            if keep_whole_table:
                keep_row = row_index < len(normalized) - 1
            else:
                keep_row = row_index == 0 or row_index == len(normalized) - 2
            paragraph.paragraph_format.keep_with_next = keep_row
            add_inline_runs(
                paragraph,
                value,
                size=TABLE_SIZE,
                bold=(row_index == 0),
            )
    # Top border is applied by the first row, bottom by the last row.
    for col_index in range(cols):
        set_cell_border(table.cell(0, col_index), top={"sz": 12, "color": "000000"})


def add_image(document, relative_path: str, caption: str) -> None:
    image_path = (SOURCE.parent / relative_path).resolve()
    if not image_path.exists():
        raise FileNotFoundError(f"missing report image: {image_path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(5.2))
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.first_line_indent = Pt(0)
    add_inline_runs(caption_paragraph, caption, size=CAPTION_SIZE, color=GRAY)


# ---------------------------------------------------------------------------
# Cover and front matter
# ---------------------------------------------------------------------------


def build_cover(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Cm(0))
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table)
    set_table_width(table, Cm(21))
    clear_table_borders(table)
    table.columns[0].width = Cm(21)
    cell = table.cell(0, 0)
    cell.width = Cm(21)
    set_cell_border(cell, top=None, left=None, bottom=None, right=None)
    set_cell_margins(
        cell,
        top=Cm(2.2),
        bottom=Cm(2.2),
        left=Cm(3.0),
        right=Cm(2.5),
    )
    row = table.rows[0]
    row.height = Cm(29.4)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    first = cell.paragraphs[0]
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(first, before=18, after=0)
    run = first.add_run("华东师范大学")
    set_run_font(run, size=22, bold=False, color=BLACK)
    add_char_spacing(run, 40)

    title = cell.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(title, before=30, after=0)
    run = title.add_run("基于 MongoDB 的")
    set_run_font(run, cn=HEAD_CN, size=30, bold=True, color=BLACK)

    title2 = cell.add_paragraph()
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title2.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(title2, before=0, after=10)
    run = title2.add_run("在线问卷系统")
    set_run_font(run, cn=HEAD_CN, size=30, bold=True, color=BLACK)

    subtitle = cell.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(subtitle, before=4, after=0)
    run = subtitle.add_run("第二阶段项目完成报告")
    set_run_font(run, size=18, color=BLACK)

    studio = cell.add_paragraph()
    studio.alignment = WD_ALIGN_PARAGRAPH.CENTER
    studio.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(studio, before=8, after=0)
    run = studio.add_run("问研 Survey Studio")
    set_run_font(run, size=12, color=GRAY)

    spacer = cell.add_paragraph()
    spacer.paragraph_format.first_line_indent = Pt(0)
    spacer.paragraph_format.space_before = Pt(32)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0

    meta = cell.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    set_table_layout_fixed(meta)
    set_table_width(meta, Cm(11.0))
    clear_table_borders(meta)
    meta.columns[0].width = Cm(3.6)
    meta.columns[1].width = Cm(7.4)
    entries = [
        ("课程名称：", "数据库课程大作业"),
        ("项目名称：", "基于 MongoDB 的在线问卷系统"),
        ("作者：", "王子腾"),
        ("学号：", "10245501473"),
    ]
    for index, (label, value) in enumerate(entries):
        row = meta.rows[index]
        row.height = Cm(0.85)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        label_cell = meta.cell(index, 0)
        value_cell = meta.cell(index, 1)
        label_cell.width = Cm(3.6)
        value_cell.width = Cm(7.4)
        for current in (label_cell, value_cell):
            set_cell_margins(current, top=Cm(0.06), bottom=Cm(0.06), left=Cm(0.1), right=Cm(0.1))
            set_cell_border(current, top=None, left=None, right=None, bottom=None)
        set_cell_border(value_cell, bottom={"sz": 6, "color": "000000"})
        for current, text, bold in ((label_cell, label, False), (value_cell, value, False)):
            p = current.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = LINE
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, size=11, bold=bold, color=BLACK)

    footer = cell.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(footer, before=20, after=0)
    run = footer.add_run("2026 年 10 月")
    set_run_font(run, size=10.5, color=LIGHT_GRAY)


def add_char_spacing(run, value: int) -> None:
    rpr = run._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(value))
    rpr.append(spacing)


def configure_section(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.3)
    section.footer_distance = Cm(1.4)


def add_header_footer(section, *, roman: bool) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.first_line_indent = Pt(0)
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("问研 Survey Studio · 第二阶段项目完成报告")
    set_run_font(run, size=9, color=GRAY)
    add_bottom_border(header_p, size=4, color="D9D9D9")

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.first_line_indent = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    run = footer_p.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    add_field(footer_p, "PAGE \\* ROMAN \\* MERGEFORMAT" if roman else "PAGE \\* arabic \\* MERGEFORMAT", "i")
    run = footer_p.add_run(" 页")
    set_run_font(run, size=9, color=GRAY)


def build_toc_page(document: Document) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    add_header_footer(section, roman=True)
    set_page_numbering(section, "upperRoman")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(title, before=6, after=18)
    run = title.add_run("目  录")
    set_run_font(run, cn=HEAD_CN, size=18, bold=True, color=BLACK)

    toc_p = document.add_paragraph()
    toc_p.paragraph_format.first_line_indent = Pt(0)
    toc_p.paragraph_format.line_spacing = LINE
    toc_p.paragraph_format.space_after = Pt(6)
    add_field(toc_p, 'TOC \\o "1-2" \\h \\z \\u', "目录将在打开文档时自动生成")

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_spacing(note, before=8, after=0)
    run = note.add_run("提示：如目录页码未自动刷新，请在目录上右键选择“更新域”。")
    set_run_font(run, size=9, color=GRAY, italic=True)


def build_body_section(document: Document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    add_header_footer(section, roman=False)
    set_page_numbering(section, "decimal")
    return section


# ---------------------------------------------------------------------------
# Markdown renderer
# ---------------------------------------------------------------------------


def render_markdown(document: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line.rstrip())
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("![") and "](" in stripped:
            match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", stripped)
            if match:
                add_image(document, match.group(2), match.group(1))
            index += 1
            continue

        if stripped.startswith("# "):
            text = stripped[2:].strip()
            if text.startswith("项目完成报告"):
                index += 1
                continue
            add_heading(document, text, 1)
            index += 1
            continue

        if stripped.startswith("## "):
            add_heading(document, stripped[3:].strip(), 2)
            index += 1
            continue

        if stripped.startswith("### "):
            add_heading(document, stripped[4:].strip(), 3)
            index += 1
            continue

        if stripped == "---":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            set_paragraph_spacing(paragraph, before=4, after=4)
            run = paragraph.add_run("— — —")
            set_run_font(run, size=9, color=LIGHT_GRAY)
            index += 1
            continue

        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                if not is_separator(lines[index]):
                    rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        if stripped.startswith("> "):
            add_blockquote(document, stripped[2:])
            index += 1
            continue

        if stripped.startswith("- "):
            add_list_item(document, stripped[2:], ordered=False)
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            add_list_item(document, f"{numbered.group(1)}. {numbered.group(2)}", ordered=True)
            index += 1
            continue

        if stripped.startswith("**表 ") and stripped.endswith("**"):
            add_body_paragraph(document, stripped, indent=False, caption=True)
        elif stripped.startswith("关键词："):
            add_body_paragraph(document, stripped, indent=False)
        else:
            paragraph = add_body_paragraph(document, stripped, indent=True)
            next_line = ""
            next_offset = index
            for offset, candidate in enumerate(lines[index + 1 :], start=index + 1):
                if candidate.strip():
                    next_line = candidate.strip()
                    next_offset = offset
                    break
            keep = next_line.startswith("```")
            if next_line.startswith(("**表 ", "|")):
                keep = count_table_rows(lines, next_offset) <= 6
            if keep:
                paragraph.paragraph_format.keep_with_next = True
        index += 1


# ---------------------------------------------------------------------------
# Word field update (TOC + page numbers)
# ---------------------------------------------------------------------------


def update_fields_with_word(path: Path) -> None:
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        print(f"Word field update skipped: {exc}")
        return
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(path), ReadOnly=False)
        doc.Fields.Update()
        for index in range(1, doc.TablesOfContents.Count + 1):
            doc.TablesOfContents(index).Update()
        doc.Repaginate()
        doc.Save()
        print("Word fields updated (TOC and page numbers)")
    except Exception as exc:
        print(f"Word field update skipped: {exc}")
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


# ---------------------------------------------------------------------------
# Post-Word normalization
# ---------------------------------------------------------------------------


def normalize_heading_styles(path: Path) -> None:
    """Keep standard Heading1-3 style IDs after Word has saved the document.

    Chinese Word may rewrite built-in heading style IDs to ``1``/``21``/``31``
    and lowercase names. The document still renders correctly, but the style
    IDs make TOC/style audits unreliable. Normalize them back without touching
    the generated TOC field content.
    """

    style_map = {
        "1": ("Heading1", "Heading 1"),
        "21": ("Heading2", "Heading 2"),
        "31": ("Heading3", "Heading 3"),
    }
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": w_ns}

    def qnw(tag: str) -> str:
        return f"{{{w_ns}}}{tag}"

    with zipfile.ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}

    styles = etree.fromstring(entries["word/styles.xml"])
    for style in styles.xpath("//w:style", namespaces=ns):
        sid = style.get(qnw("styleId"))
        if sid not in style_map:
            continue
        new_id, new_name = style_map[sid]
        style.set(qnw("styleId"), new_id)
        name = style.find("w:name", ns)
        if name is not None:
            name.set(qnw("val"), new_name)
    entries["word/styles.xml"] = etree.tostring(
        styles, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    document = etree.fromstring(entries["word/document.xml"])
    for style in document.xpath("//w:pStyle", namespaces=ns):
        val = style.get(qnw("val"))
        if val in style_map:
            style.set(qnw("val"), style_map[val][0])
    entries["word/document.xml"] = etree.tostring(
        document, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    settings = etree.fromstring(entries["word/settings.xml"])
    if not settings.xpath("//w:updateFields", namespaces=ns):
        update = etree.SubElement(settings, qnw("updateFields"))
        update.set(qnw("val"), "true")
    entries["word/settings.xml"] = etree.tostring(
        settings, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, data in entries.items():
            target.writestr(name, data)


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


def build() -> int:
    document = Document()
    document.core_properties.title = "基于 MongoDB 的在线问卷系统 · 第二阶段项目完成报告"
    document.core_properties.subject = "数据库课程大作业第二阶段"
    document.core_properties.author = "王子腾"
    document.core_properties.comments = "问研 Survey Studio 第二阶段完成报告"
    configure_styles(document)
    add_update_fields_setting(document)

    build_cover(document)
    build_toc_page(document)
    build_body_section(document)
    render_markdown(document)

    document.save(TARGET)
    print(f"wrote {TARGET} ({TARGET.stat().st_size} bytes)")
    update_fields_with_word(TARGET)
    normalize_heading_styles(TARGET)
    print(f"final {TARGET} ({TARGET.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(build())
